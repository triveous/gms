import json
import os
import sys

import docx
import frappe
import google.generativeai as genai
import openpyxl
from frappe.model.document import Document


def extract_text_from_docx(file_path):
	"""Extracts all text from a .docx file."""
	try:
		doc = docx.Document(file_path)
		full_text = []
		for para in doc.paragraphs:
			full_text.append(para.text)
		for table in doc.tables:
			for row in table.rows:
				for cell in row.cells:
					full_text.append(cell.text)
		return "\n".join(full_text)
	except Exception as e:
		print(f"Error reading docx file {file_path}: {e}")
		sys.exit(1)


def extract_text_from_xlsx(file_path):
	"""Extracts all text from an .xlsx file, preserving some sheet/row structure."""
	try:
		workbook = openpyxl.load_workbook(file_path)
		full_text = []
		for sheet_name in workbook.sheetnames:
			full_text.append(f"--- Sheet: {sheet_name} ---")
			sheet = workbook[sheet_name]
			for row in sheet.iter_rows():
				row_text = [str(cell.value) if cell.value is not None else "" for cell in row]
				full_text.append("\t".join(row_text))
		return "\n".join(full_text)
	except Exception as e:
		print(f"Error reading xlsx file {file_path}: {e}")
		sys.exit(1)


def get_structured_data_from_gemini(text_content, api_key):
	"""Sends the extracted text to the Gemini model and asks for structured JSON output."""
	try:
		genai.configure(api_key=api_key)
		model = genai.GenerativeModel("gemini-2.0-flash")

		prompt = f"""
        You are an expert data extraction AI. Your task is to analyze the following text content extracted from a document and convert it into a structured JSON format. The structure of the original document is unknown.

        Follow these instructions carefully:
        1. Identify the main entities, relationships, and key data points.
        2. If you see tabular data, represent it as an array of JSON objects, where each object is a row. Use the table headers as keys if available; otherwise, use generic but descriptive keys.
        3. If you see key-value pairs (e.g., "Name: John Doe"), represent them as a JSON object.
        4. Group related information together under logical keys.
        5. Clean the data: remove unnecessary whitespace and artifacts.
        6. The final output MUST be a single, valid JSON object. Do not include any explanatory text, comments, or markdown code fences (like ```json) before or after the JSON.
        7. Don't include extra object under table name. Under table name, directly have the array of objects(row).
        8. Don't include sheet name, there should be single JSON output only. In that json output, directly have the key-value pairs where key is the name of the table(in each sheet we have table name, use that only) and value is the array of objects(rows).
        8. The field name should be in snake_case everytime.

        Here is the document text:
        --- DOCUMENT TEXT START ---
        {text_content}
        --- DOCUMENT TEXT END ---

        Please provide the structured JSON output and the field name should be in snake_case everytime now.
        """

		print("\nSending request to Gemini API... (This may take a moment)")
		response = model.generate_content(prompt)
		print("\nReceived response from Gemini API. Processing...", response.text)
		# Clean up potential markdown formatting from the response
		json_response = response.text.strip().replace("```json", "").replace("```", "").strip()

		return json.loads(json_response)

	except Exception as e:
		print(f"An error occurred with the Gemini API call: {e}")
		if "response" in locals() and hasattr(response, "prompt_feedback"):
			print(f"Prompt Feedback: {response.prompt_feedback}")
		sys.exit(1)


def extract_file_text(file_path):
	"""Auto-detect extension and extract."""
	ext = os.path.splitext(file_path)[1].lower()

	if ext == ".docx":
		return extract_text_from_docx(file_path)

	if ext == ".xlsx":
		return extract_text_from_xlsx(file_path)

	raise Exception(f"Unsupported file type: {ext}")


def run_extraction_pipeline(file_path):
	"""Master reusable function used by any DocType."""
	api_key = frappe.conf.get("google_api_key")
	text = extract_file_text(file_path)
	if not text.strip():
		raise Exception("Empty document")

	return get_structured_data_from_gemini(text, api_key=api_key)
